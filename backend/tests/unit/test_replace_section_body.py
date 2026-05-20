import pytest
import io
from docx import Document
from app.services.docx_template_renderer import DocxTemplateRenderer
from app.services.resume_generator_service import ResumeGeneratorService

def test_docx_template_renderer_replace_section_content():
    # 1. Create a dummy document
    doc = Document()
    doc.add_heading("Header Info", level=1)
    doc.add_paragraph("Some top text")
    
    # Target heading section
    doc.add_paragraph("WORK EXPERIENCE")
    doc.add_paragraph("[Organisation 1]")
    doc.add_paragraph("[Responsibilities]")
    doc.add_paragraph("[Bullet 1]")
    
    # Boundary (Next heading or uppercase block)
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("[University]")
    
    # 2. Run _replace_section_content
    renderer = DocxTemplateRenderer()
    renderer._replace_section_content(doc, "WORK EXPERIENCE", "{{ _['work_experience'] }}")
    
    # 3. Assertions
    paragraphs = [p.text for p in doc.paragraphs]
    assert "WORK EXPERIENCE" in paragraphs
    assert "{{ _['work_experience'] }}" in paragraphs
    assert "[Organisation 1]" not in paragraphs
    assert "[Responsibilities]" not in paragraphs
    assert "[Bullet 1]" not in paragraphs
    assert "EDUCATION" in paragraphs
    assert "[University]" in paragraphs

def test_resume_generator_service_replace_section_body():
    # 1. Create a dummy document
    doc = Document()
    doc.add_heading("WORK EXPERIENCE", level=1)
    doc.add_paragraph("[Organisation 1]")
    doc.add_paragraph("[Responsibilities]")
    
    # Add a table to act as a boundary
    tbl = doc.add_table(rows=1, cols=1)
    tbl.cell(0, 0).text = "Table Cell Content"
    
    doc.add_paragraph("Some text after table")
    
    # 2. Setup locator field manifest and render context
    field_manifest = [
        {
            "fieldname": "work_experience",
            "render_locator": {
                "strategy": "replace_section_body",
                "heading": "WORK EXPERIENCE"
            }
        }
    ]
    render_context = {
        "work_experience": "Fabulous job experience details"
    }
    
    # 3. Run apply_render_locators
    service = ResumeGeneratorService()
    service.apply_render_locators(doc, field_manifest, render_context)
    
    # 4. Assertions
    paragraphs = [p.text for p in doc.paragraphs]
    assert "{{ _['work_experience'] }}" in paragraphs
    assert "[Organisation 1]" not in paragraphs
    assert "[Responsibilities]" not in paragraphs
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(0, 0).text == "Table Cell Content"

def test_resume_generator_service_replace_bullet_list_under_heading():
    doc = Document()
    doc.add_heading("Key skills", level=1)
    doc.add_paragraph("[Type text]")
    doc.add_paragraph("«Type text»")
    field_manifest = [{
        "fieldname": "skills",
        "render_locator": {"strategy": "replace_bullet_list_under_heading", "heading": "Key skills"}
    }]
    render_context = {"skills": ["PyTorch", "OpenAI API"]}
    service = ResumeGeneratorService()
    service.apply_render_locators(doc, field_manifest, render_context)
    paragraphs = [p.text for p in doc.paragraphs]
    assert "[Type text]" not in paragraphs
    assert "«Type text»" not in paragraphs
    assert "• PyTorch" in paragraphs
    assert "• OpenAI API" in paragraphs

def test_prepare_document_markers_preserves_unresolved_marker_when_value_empty():
    doc = Document()
    doc.add_paragraph("Recruiting experts in «EmployeeJobTitle»")
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    service = ResumeGeneratorService()
    out = service.prepare_document_markers(
        stream,
        field_list=["employee_job_title"],
        field_manifest=[],
        resume_data={"employee_job_title": ""},
    )
    out_doc = Document(out)
    texts = [p.text for p in out_doc.paragraphs]
    assert any("«EmployeeJobTitle»" in t for t in texts)
