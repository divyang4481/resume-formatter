import sys
import os
sys.path.append(os.path.abspath('backend'))
from docx import Document
from app.services.document_marker_locator import DocumentMarkerLocator

doc = Document(r'SampleData\templates\UK Worldwide London.docx')
locator = DocumentMarkerLocator()

def is_heading_match(p_text, target_heading):
    import re
    if not p_text or not target_heading:
        return False
    norm_p = re.sub(r'[^a-zA-Z0-9]', '', p_text).lower()
    norm_t = re.sub(r'[^a-zA-Z0-9]', '', target_heading).lower()
    return norm_p == norm_t or (len(norm_t) > 3 and norm_t in norm_p and len(norm_p) < len(norm_t) + 10)

def is_heading_boundary(p):
    text = p.text.strip()
    if not text:
        return False
    style_name = p.style.name.lower() if p.style and p.style.name else ""
    if "heading" in style_name or style_name.startswith("h") and any(style_name.endswith(str(i)) for i in range(1, 7)):
        return True
    is_bold = any(run.bold for run in p.runs)
    if is_bold and len(text) < 60 and not text.startswith("["):
        return True
    if text.isupper() and len(text) < 60:
        return True
    return False

found = False
for i, p in enumerate(doc.paragraphs):
    if is_heading_match(p.text.strip(), "Skills"):
        found = True
        print(f"FOUND SKILLS at {i}: '{p.text}'")
        for j in range(i+1, min(i+10, len(doc.paragraphs))):
            next_p = doc.paragraphs[j]
            print(f"  Paragraph {j}: '{next_p.text}' -> Boundary: {is_heading_boundary(next_p)}, Upper: {next_p.text.strip().isupper()}, Len: {len(next_p.text.strip())}, Style: {next_p.style.name}")
        break

