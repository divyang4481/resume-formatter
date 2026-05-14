from docx import Document
import sys

doc = Document(sys.argv[1])
print("--- MAIN BODY ---")
for p in doc.paragraphs:
    if p.text.strip():
        print(f"P: {p.text}")

print("\n--- TABLES ---")
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            if cell.text.strip():
                print(f"CELL: {cell.text}")

print("\n--- HEADERS ---")
for section in doc.sections:
    if section.header:
        for p in section.header.paragraphs:
            if p.text.strip():
                print(f"H: {p.text}")
