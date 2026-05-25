import sys, os, json, logging
sys.path.append(os.path.abspath('backend'))
from docx import Document
from docxtpl import DocxTemplate
from app.services.document_marker_locator import DocumentMarkerLocator
from app.services.rich_text_renderer import RichTextRenderer
import io

logging.basicConfig(level=logging.INFO)

doc_path = r'SampleData\templates\UK Worldwide London.docx'
doc = Document(doc_path)
locator = DocumentMarkerLocator()

manifest = {
    "fields": [
        {
            "fieldname": "skills",
            "marker_text": "[Type text]",
            "field_type": "list",
            "render_locator": {
                "strategy": "replace_section_body",
                "heading": "Skills"
            }
        },
        {
            "fieldname": "work_experience",
            "marker_text": "WORK EXPERIENCE",
            "field_type": "list",
            "render_locator": {
                "strategy": "replace_complex_block",
                "heading": "WORK EXPERIENCE"
            }
        }
    ]
}

context = {
    "skills": ["Python", "Docker", "AWS"],
    "work_experience": [
        {"job_title": "SWE", "company": "Google", "start_date": "2020", "end_date": "2022", "description": "Did stuff"}
    ]
}

# simulate format_array_complex_value
context['skills_str'] = '\n'.join(f"- {x}" for x in context['skills'])
context['work_experience_str'] = "SWE | Google (2020 - 2022)\n  Did stuff"

locator.apply_render_locators(doc, manifest, context)

# save temp
temp_io = io.BytesIO()
doc.save(temp_io)
temp_io.seek(0)

# docxtpl
tpl = DocxTemplate(temp_io)
rt_renderer = RichTextRenderer()
render_context = rt_renderer.apply_rendering_actions(context)
render_context_with_scope = {**render_context, "_": render_context}
tpl.render(render_context_with_scope)

out_io = io.BytesIO()
tpl.save(out_io)
out_io.seek(0)
final_doc = Document(out_io)

for i, p in enumerate(final_doc.paragraphs):
    if p.text.strip():
        print(f"Para {i}: {repr(p.text)}")
