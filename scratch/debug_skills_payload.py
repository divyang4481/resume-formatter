import sys
import os
sys.path.append(os.path.abspath('backend'))
import json
from docx import Document
from app.services.document_marker_locator import DocumentMarkerLocator

import logging
logging.basicConfig(level=logging.INFO)

doc = Document(r'SampleData\templates\UK Worldwide London.docx')
with open(r'scratch\latest_job_data.json', 'r', encoding='utf-8') as f:
    job_data = json.load(f)

# Maybe the job data is the manifest itself if it has 'fields'
if 'fields' in job_data:
    manifest = job_data
elif 'template_manifest' in job_data and 'fields' in job_data['template_manifest']:
    manifest = job_data['template_manifest']
elif 'job' in job_data and 'template_manifest' in job_data['job']:
    manifest = job_data['job']['template_manifest']
else:
    print("Could not find manifest fields!")
    sys.exit(1)

print("Manifest fields count:", len(manifest.get('fields', [])))

mock_render_context = {}
for f in manifest.get("fields", []):
    fieldname = f.get("fieldname") or f.get("canonical_fieldname")
    if fieldname:
        f["fieldname"] = fieldname # ensure it's there
        if f.get("field_type") == "list":
            mock_render_context[fieldname] = ["Skill 1", "Skill 2"]
        else:
            mock_render_context[fieldname] = "test value"
            mock_render_context[fieldname + "_str"] = "test value str"

locator = DocumentMarkerLocator()
locator.apply_render_locators(doc, manifest, mock_render_context)

found = False
for p in doc.paragraphs:
    if "WORK EXPERIENCE" in p.text:
        found = True
        print("FOUND WORK EXPERIENCE in modified doc")
if not found:
    print("WORK EXPERIENCE IS GONE!")

for p in doc.paragraphs:
    if "Skill 1" in p.text:
        print("FOUND Skill 1 in modified doc")

